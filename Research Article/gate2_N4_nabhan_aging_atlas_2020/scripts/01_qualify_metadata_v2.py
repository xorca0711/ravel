"""M0: identify source units, coverage and matrix scale; no gene-level contrasts."""
from pathlib import Path
import argparse,json,platform,sys
import h5py,numpy as np,pandas as pd,scipy
from openpyxl import load_workbook
from nb5_io import read_obs,read_matrix,read_genes,write_json,write_tsv

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--manifest',required=True);ap.add_argument('--output',required=True);args=ap.parse_args()
    out=Path(args.output);out.mkdir(exist_ok=True)
    manifest=json.loads(Path(args.manifest).read_text())
    cells=[];audits=[];schema=[];sheets=[];published=[]
    for record in manifest['files']:
        path=Path(record['path'])
        if path.suffix=='.h5ad':
            dataset=path.stem
            with h5py.File(path,'r') as f:
                obs=read_obs(f)
                for required in ['age','sex','mouse_id','tissue','cell_ontology_class','method']:
                    if required not in obs:raise ValueError(f'{dataset}: missing {required}')
                obs['dataset']=dataset
                obs['age_months']=pd.to_numeric(obs.age.astype(str).str.replace('m','',regex=False),errors='raise')
                columns=['dataset','cell_id','mouse_id','age','age_months','sex','tissue','method','cell_ontology_class']
                columns+=[c for c in ['subtissue','plate','well','louvain','leiden','cluster_names','cluster_number','n_counts','n_genes'] if c in obs]
                cells.append(obs[columns])
                schema.append({'dataset':dataset,'obs_fields':list(obs.columns),'root_keys':list(f.keys()),'uns_keys':list(f['uns'].keys()) if 'uns' in f else [],'n_cells':len(obs)})
                for key in ['X','raw.X']:
                    if key not in f:continue
                    x=read_matrix(f,key);v=x.data
                    finite=bool(np.isfinite(v).all());nonnegative=bool((v>=0).all())
                    integral=bool(finite and np.allclose(v,np.rint(v),rtol=0,atol=1e-6))
                    sums=np.asarray(x.sum(axis=1),dtype=float).ravel()
                    ncounts=pd.to_numeric(obs.get('n_counts',pd.Series(np.nan,index=obs.index)),errors='coerce').to_numpy()
                    match=bool(np.isfinite(ncounts).all() and np.allclose(sums,ncounts,rtol=1e-5,atol=1))
                    genes=read_genes(f,'raw.var' if key=='raw.X' else 'var')
                    audits.append(dict(dataset=dataset,layer=key,n_cells=x.shape[0],n_genes=x.shape[1],nonzero=x.nnz,min_nonzero=float(v.min()) if len(v) else None,max_nonzero=float(v.max()) if len(v) else None,finite=finite,nonnegative=nonnegative,all_integer=integral,row_sum_matches_deposited_n_counts=match,median_row_sum=float(np.median(sums)),duplicate_gene_ids=int(pd.Series(genes).duplicated().sum()),count_scale_supported=bool(integral and nonnegative and match)))
                    del x
        elif path.suffix=='.xlsx':
            wb=load_workbook(path,read_only=True,data_only=True)
            for ws in wb:
                first=next(ws.values,())
                sheets.append({'file':path.name,'sheet':ws.title,'rows':ws.max_row,'columns':ws.max_column,'first_row':[str(x) if x is not None else '' for x in first]})
            if path.name in {'supplement_table_1.xlsx','supplement_table_2.xlsx'}:
                # Decode only explicitly merged cells; never infer values across blank rows.
                full=load_workbook(path,read_only=False,data_only=True)
                ws=full['a']; records=[list(row) for row in ws.values]
                for span in ws.merged_cells.ranges:
                    value=ws.cell(span.min_row,span.min_col).value
                    for rr in range(span.min_row,span.max_row+1):
                        for cc in range(span.min_col,span.max_col+1):records[rr-1][cc-1]=value
                full.close()
                table=pd.DataFrame(records[1:],columns=records[0]).dropna(how='all')
                # Totals/separators have no tissue and are not animal-tissue records.
                table=table.loc[table['Tissue'].notna()].copy()
                table=table.rename(columns={'Age':'age','Sex':'sex','Mouse ID':'mouse_id','Tissue':'tissue','Number of cells':'source_cells'})
                table['method']='facs' if '_1.' in path.name else 'droplet'
                table['source_table']=path.name
                published.append(table)
            wb.close()
    all_cells=pd.concat(cells,ignore_index=True)
    for col in ['mouse_id','age','sex','tissue','method','cell_ontology_class']:
        all_cells[col]=all_cells[col].fillna('').astype(str)
    write_tsv(out/'cell_metadata.tsv',all_cells)
    keys=['dataset','mouse_id','age','age_months','sex','tissue','method','cell_ontology_class']
    coverage=all_cells.groupby(keys,dropna=False,observed=True).size().rename('cells').reset_index()
    write_tsv(out/'coverage.tsv',coverage)
    animal_keys=['dataset','mouse_id','age','age_months','sex','tissue','method']
    animal=all_cells.groupby(animal_keys,observed=True,dropna=False).size().rename('cells').reset_index()
    write_tsv(out/'animals.tsv',animal)
    conflicts=[]
    for mouse,frame in all_cells.groupby('mouse_id',observed=True):
        if not mouse or frame.age.nunique()!=1 or frame.sex.nunique()!=1:
            conflicts.append({'mouse_id':mouse,'ages':sorted(frame.age.unique()),'sexes':sorted(frame.sex.unique())})
    source=pd.concat(published,ignore_index=True)
    for col in ['mouse_id','age','sex','tissue','method']:source[col]=source[col].fillna('').astype(str)
    write_tsv(out/'published_coverage.tsv',source)
    join=['mouse_id','age','sex','tissue','method']
    check=animal.merge(source,on=join,how='left',validate='many_to_one',indicator=True)
    check['cell_difference']=check.cells-check.source_cells
    write_tsv(out/'source_concordance.tsv',check)
    design=animal.groupby(['dataset','age_months','sex'],observed=True).agg(animals=('mouse_id','nunique'),cells=('cells','sum')).reset_index()
    write_tsv(out/'design.tsv',design)
    write_tsv(out/'layer_audit.tsv',pd.DataFrame(audits))
    write_json(out/'sheet_catalog.json',sheets)
    write_json(out/'inventory.json',{'source_manifest':args.manifest,'author_code_commit':manifest['author_code_commit'],'datasets':schema,'total_observations':len(all_cells),'unique_deposited_mouse_ids':int(all_cells.mouse_id.nunique()),'mouse_identity_conflicts':conflicts,'missing_identity_cells':int((all_cells.mouse_id=='').sum()),'source_unmatched_rows':int((check['_merge']!='both').sum()),'source_count_mismatch_rows':int((check.cell_difference.fillna(0)!=0).sum()),'interpretation':'Metadata and scale qualification only. Sample counts do not justify causal or confirmatory inference. Raw-versus-normalized classification requires source provenance as well as numerical checks.','environment':{'python':sys.version.split()[0],'numpy':np.__version__,'pandas':pd.__version__,'scipy':scipy.__version__,'h5py':h5py.__version__}})
    print(json.dumps({'observations':len(all_cells),'datasets':len(cells),'identity_conflicts':len(conflicts),'count_scale_supported_layers':sum(a['count_scale_supported'] for a in audits)}))

if __name__=='__main__':main()
