"""Bounded exposed Nb5 source reproduction and mouse-level description.

No count model, cell-level hypothesis test, state discovery or causal inference.
"""
from pathlib import Path
import argparse,json,re,sys
from collections import Counter
import h5py,numpy as np,pandas as pd
from scipy.stats import hypergeom
from openpyxl import load_workbook
from nb5_io import read_obs,read_matrix,read_genes,write_json,write_tsv

MARKERS=['Cdkn2a','Cdkn1a','Lmnb1','Il1b','Tnf','Cd38','H2-D1','H2-K1','B2m','Fth1','Oasl2','Oas1a','Ifit3','Rtp4','Bst2','Stat1','Irf7','Ifitm3','Usp18','Ifi204','Ifit2','P2ry12','Tmem119','Apoe','Lpl','Cst7','Clec7a','Itgal','Itga1','Msrb1','Spex1']
CASE='facs.Brain_Myeloid.clustered_diversity'
KEY=['dataset','mouse_id','age_months','sex','unit_kind']

def unit_kind(s):return 'pooled_mouse_label' if '/' in str(s) else 'deposited_mouse'
def table(w,s,skip=0):
    rows=w[s].values
    for _ in range(skip):next(rows)
    cols=list(next(rows));cols=[x if x is not None else 'gene' for x in cols]
    return pd.DataFrame(rows,columns=cols).dropna(how='all')

def summarize_composition(obs,dataset):
    rows=[];labels=sorted(obs.cell_ontology_class.unique())
    for (mouse,age,sex),g in obs.groupby(['mouse_id','age_months','sex'],observed=True):
        counts=g.cell_ontology_class.value_counts()
        for label in labels:
            rows.append(dict(dataset=dataset,mouse_id=mouse,age_months=int(age),sex=sex,unit_kind=unit_kind(mouse),cell_type=label,cells=int(counts.get(label,0)),total_cells=len(g),fraction=float(counts.get(label,0)/len(g))))
    return rows

def expression(obs,x,genes,dataset,scale):
    rows=[];available=[g for g in MARKERS if g in genes];ix=[int(np.flatnonzero(genes==g)[0]) for g in available]
    y=x[:,ix].toarray().astype(float)
    for (mouse,age,sex),indices in obs.groupby(['mouse_id','age_months','sex'],observed=True).groups.items():
        indices=np.asarray(list(indices));sub=obs.loc[indices]
        groups={'ALL':indices}
        groups.update({k:np.asarray(list(v)) for k,v in sub.groupby('cell_ontology_class',observed=True).groups.items()})
        for label,idx in groups.items():
            z=y[idx];positive=(z>0).sum(axis=0);sums=z.sum(axis=0)
            for j,gene in enumerate(available):
                rows.append(dict(dataset=dataset,mouse_id=mouse,age_months=int(age),sex=sex,unit_kind=unit_kind(mouse),cell_type=label,gene=gene,cells=len(idx),positive_cells=int(positive[j]),detected_fraction=float(positive[j]/len(idx)),mean_expression=float(sums[j]/len(idx)),positive_mean=float(sums[j]/positive[j]) if positive[j] else np.nan,scale=scale))
    return rows,available

def contrasts(expr):
    rows=[]
    for (dataset,ct,gene,scale),g in expr.groupby(['dataset','cell_type','gene','scale'],observed=True):
        g=g[g.unit_kind=='deposited_mouse']
        for sex_scope in ['all_observed_sexes','male_only']:
            f=g if sex_scope.startswith('all') else g[g.sex=='male']
            a=f[f.age_months==3];b=f[f.age_months==24]
            if a.empty or b.empty:continue
            for endpoint in ['detected_fraction','mean_expression']:
                rows.append(dict(dataset=dataset,cell_type=ct,gene=gene,scale=scale,sex_scope=sex_scope,endpoint=endpoint,n_3m=len(a),n_24m=len(b),mean_3m=float(a[endpoint].mean()),mean_24m=float(b[endpoint].mean()),difference_24_minus_3=float(b[endpoint].mean()-a[endpoint].mean()),young_sexes=';'.join(sorted(a.sex.unique())),old_sexes=';'.join(sorted(b.sex.unique()))))
    return pd.DataFrame(rows)

def accounting(expr,comp):
    rows=[]
    for dataset in ['Lung_facs','Lung_droplet']:
        e=expr[(expr.dataset==dataset)&(expr.cell_type!='ALL')&(expr.unit_kind=='deposited_mouse')]
        c=comp[(comp.dataset==dataset)&(comp.unit_kind=='deposited_mouse')]
        # Fixed standard: equal cell-type weights among types seen in every retained mouse.
        units=set(c.mouse_id);support=c[c.cells>0].groupby('cell_type').mouse_id.nunique()
        common=support[support==len(units)].index.tolist()
        for (mouse,age,sex,gene),g in e.groupby(['mouse_id','age_months','sex','gene'],observed=True):
            allrow=expr[(expr.dataset==dataset)&(expr.mouse_id==mouse)&(expr.cell_type=='ALL')&(expr.gene==gene)].iloc[0]
            s=g[g.cell_type.isin(common)]
            rows.append(dict(dataset=dataset,mouse_id=mouse,age_months=int(age),sex=sex,gene=gene,observed_mean=float(allrow.mean_expression),equal_type_mean=float(s.mean_expression.mean()) if len(s)==len(common) and common else np.nan,common_types=';'.join(common),n_common_types=len(common),covered_fraction=float(c[(c.mouse_id==mouse)&c.cell_type.isin(common)].fraction.sum())))
    return pd.DataFrame(rows)

def repertoire(raw,out):
    w=load_workbook(raw/'supplement_table_9.xlsx',read_only=True,data_only=True);meta=table(w,'metadata');items=[];audit=[]
    for age in [3,18,24]:
        f=table(w,f'cell_data_tracer_{age}m');key='cell_name_3m' if age==3 else 'cell_name_18m_24m'
        m=meta[meta.age==f'{age}m'][[key,'mouse.id','tissue']].dropna(subset=[key]).rename(columns={key:'cell_name','mouse.id':'mouse_id'})
        if m.cell_name.duplicated().any():raise ValueError('Ambiguous source repertoire cell mapping')
        if f.cell_name.duplicated().any():raise ValueError('Repeated T-cell source row')
        d=f[['cell_name','clonal_group','group_size']].copy();d['any_productive']=f[['A_productive','B_productive']].notna().any(axis=1);d['age_months']=age
        d['source_in_clone']=pd.to_numeric(d.group_size,errors='coerce').fillna(1)>1
        d=d.merge(m,on='cell_name',how='left',validate='one_to_one',indicator=True)
        audit.append(dict(age_months=age,source_rows=len(d),source_clone_cells=int(d.source_in_clone.sum()),pooled_fraction=float(d.source_in_clone.mean()),mapped_rows=int((d['_merge']=='both').sum()),unmapped_rows=int((d['_merge']!='both').sum()),productive_rows=int(d.any_productive.sum())))
        items.append(d.drop(columns='_merge'))
    w.close();cells=pd.concat(items,ignore_index=True);write_tsv(out/'repertoire_cells.tsv',cells)
    write_tsv(out/'repertoire_source.tsv',pd.DataFrame(audit));summaries=[]
    # Exact common-depth expected fraction, conditional on observed within-mouse clone labels.
    mapped=cells.dropna(subset=['mouse_id']);counts=mapped.groupby(['age_months','mouse_id']).size();depth=int(counts.min()) if len(counts) else 0
    for (age,mouse),g in mapped.groupby(['age_months','mouse_id'],observed=True):
        labels=g.clonal_group.copy().astype('string')
        labels=labels.fillna(pd.Series(['singleton_'+str(i) for i in g.index],index=g.index))
        sizes=labels.value_counts().to_numpy();n=len(g)
        expected=sum(k/n*(1-hypergeom.pmf(0,n-1,k-1,depth-1)) for k in sizes) if depth>=2 else np.nan
        summaries.append(dict(age_months=int(age),mouse_id=mouse,reconstructed_cells=n,source_clone_cells=int(g.source_in_clone.sum()),source_clone_fraction=float(g.source_in_clone.mean()),within_mouse_clone_fraction=float(sizes[sizes>1].sum()/n),common_depth=depth,expected_fraction_at_common_depth=float(expected),tissues=';'.join(sorted(g.tissue.dropna().unique()))))
    write_tsv(out/'repertoire_mouse.tsv',pd.DataFrame(summaries))
    return {'source_rows':len(cells),'unmapped_rows':int(cells.mouse_id.isna().sum()),'common_depth':depth,'sampling_limit':'Conditional on reconstructed receptors; tissue mixture and assembly selection remain. Fixed labels are not reassembled. No immune-function inference.'}

def source_overlap(raw,out):
    w=load_workbook(raw/'supplement_table_10.xlsx',read_only=True,data_only=True)
    a=table(w,'facs.Brain_Myeloid.101214vs16',1);b=table(w,'Amit_expression_mic3tomic1',1);w.close()
    top=a.gene.dropna().astype(str).head(200).tolist();disease=b.gene.dropna().astype(str).head(200).tolist()
    rows=[]
    for i,g in enumerate(top):rows.append(dict(gene=g,source_age_rank=i+1,in_first_200_disease_rows=g in set(disease),in_any_disease_row=g in set(b.gene.dropna().astype(str))))
    write_tsv(out/'source_overlap.tsv',pd.DataFrame(rows))
    return dict(age_top_rows=len(top),disease_top_rows=len(disease),top200_intersection=len(set(top)&set(disease)),age_top200_in_full_disease_list=len(set(top)&set(b.gene.dropna().astype(str))),interpretation='Source-list arithmetic only; no enrichment test, held-out disease validation or marker selection from this overlap.')

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--manifest',required=True);ap.add_argument('--metadata',required=True);ap.add_argument('--output',required=True);args=ap.parse_args()
    out=Path(args.output);out.mkdir(exist_ok=True);manifest=json.loads(Path(args.manifest).read_text());metadata=Path(args.metadata)
    comp=[];expr=[];annotations=[];missing=[];verify=[];state_rows=[];state_expr=[];umap=[];case_identity=[]
    for record in manifest['files']:
        path=Path(record['path'])
        if path.suffix!='.h5ad':continue
        dataset=path.stem
        with h5py.File(path,'r') as f:
            obs=read_obs(f);obs['age_months']=obs.age.str.replace('m','',regex=False).astype(int)
            for (ct,fa),g in obs.groupby(['cell_ontology_class','free_annotation'],observed=True,dropna=False):annotations.append(dict(dataset=dataset,cell_type=str(ct),free_annotation=str(fa),cells=len(g)))
            if dataset==CASE:
                case_identity=obs[['mouse_id','age_months','sex']].drop_duplicates().to_dict('records')
                source_micro=obs.cell_ontology_class=='microglial cell'
                selected=obs.loc[source_micro].copy();selected['state']=np.where(selected.leiden.isin(['1','6']),'source_1_6',np.where(selected.leiden.isin(['10','12','14']),'source_10_12_14','other_source_clusters'))
                for (mouse,age,sex),g in selected.groupby(['mouse_id','age_months','sex'],observed=True):
                    for state in ['source_1_6','source_10_12_14','other_source_clusters']:state_rows.append(dict(mouse_id=mouse,age_months=int(age),sex=sex,state=state,cells=int((g.state==state).sum()),total_microglia=len(g),fraction=float((g.state==state).mean())))
                obsm=f['obsm'];coords=obsm['X_umap'][()] if isinstance(obsm,h5py.Group) else obsm[()]['X_umap']
                for i,row in selected.iterrows():umap.append(dict(cell_id=row.cell_id,mouse_id=row.mouse_id,age_months=int(row.age_months),sex=row.sex,leiden=row.leiden,state=row.state,umap1=float(coords[i,0]),umap2=float(coords[i,1])))
                # X has unknown transformation beyond nonnegative clipping; no reconstruction to counts.
                x=read_matrix(f,'X');genes=read_genes(f,'var');rr,available=expression(obs,x,genes,dataset,'source_transformed_X');state_expr+=rr
                verify.append(dict(check='case_microglial_source_count',dataset=dataset,values={str(a):int((selected.age_months==a).sum()) for a in [3,18,24]}))
                continue
            rr=summarize_composition(obs,dataset);comp+=rr
            # independent Counter arithmetic checks every retained unit/type, including zero cells
            direct=Counter(zip(obs.mouse_id,obs.cell_ontology_class))
            assert all(v['cells']==direct[(v['mouse_id'],v['cell_type'])] for v in rr)
            x=read_matrix(f,'raw.X');genes=read_genes(f,'raw.var')
            totals=np.asarray(x.sum(axis=1),dtype=float).ravel();assert np.allclose(totals,10000,rtol=1e-4,atol=0.02)
            rr,available=expression(obs,x,genes,dataset,'normalized_per_10000');expr+=rr
            missing.append(dict(dataset=dataset,requested=MARKERS,absent=[g for g in MARKERS if g not in available]))
            # Separate direct CSR traversal for one preselected primary marker over every animal.
            if 'Cdkn2a' in available:
                col=int(np.flatnonzero(genes=='Cdkn2a')[0]);direct_sums={};direct_n={};direct_pos={}
                for i,mouse in enumerate(obs.mouse_id):
                    start,end=x.indptr[i:i+2];where=x.indices[start:end]==col;v=float(x.data[start:end][where].sum())
                    direct_sums[mouse]=direct_sums.get(mouse,0)+v;direct_n[mouse]=direct_n.get(mouse,0)+1;direct_pos[mouse]=direct_pos.get(mouse,0)+(v>0)
                for row in rr:
                    if row['gene']=='Cdkn2a' and row['cell_type']=='ALL':
                        m=row['mouse_id'];assert abs(row['mean_expression']-direct_sums[m]/direct_n[m])<1e-10;assert row['positive_cells']==direct_pos[m]
                verify.append(dict(check='direct_CSR_Cdkn2a_and_Counter_composition',dataset=dataset,units=len(direct_n),passed=True))
    comp=pd.DataFrame(comp);expr=pd.DataFrame(expr)
    write_tsv(out/'composition_mouse.tsv',comp);write_tsv(out/'expression_mouse.tsv',expr);write_tsv(out/'expression_contrasts.tsv',contrasts(expr));write_tsv(out/'lung_accounting.tsv',accounting(expr,comp));write_tsv(out/'annotation_dictionary.tsv',pd.DataFrame(annotations));write_tsv(out/'microglia_state_mouse.tsv',pd.DataFrame(state_rows));write_tsv(out/'microglia_source_expression.tsv',pd.DataFrame(state_expr));write_tsv(out/'microglia_umap.tsv',pd.DataFrame(umap))
    # Equality holds in the measured linear normalized space, not after log or scaling.
    residual=np.abs(expr.mean_expression-expr.detected_fraction*expr.positive_mean.fillna(0));assert float(residual.max())<1e-8
    verify.append(dict(check='fraction_times_positive_mean_equals_unconditional_mean',rows=len(expr),maximum_absolute_residual=float(residual.max()),passed=True))
    raw=Path('raw_data/tabula_muris_senis_2020/intake_v1');rep=repertoire(raw,out);overlap=source_overlap(raw,out)
    atoms=pd.read_csv(metadata/'animals.tsv',sep='\t');atoms['unit_kind']=atoms.mouse_id.map(unit_kind)
    write_tsv(out/'design_units.tsv',atoms)
    source=pd.read_csv(metadata/'source_concordance.tsv',sep='\t');assert (source['_merge']=='both').all();write_tsv(out/'release_concordance.tsv',source)
    write_json(out/'verification.json',{'passed':True,'checks':verify,'scope':'Independent arithmetic routes on reused data; not independent biological replication.'})
    write_json(out/'summary.json',{'status':'executed_descriptive_only','case_identity':case_identity,'missing_markers':missing,'repertoire':rep,'source_overlap':overlap,'limits':['No raw count layer qualified.','No new state classifier or mixture model fitted: all-age source X transformation remains unresolved.','No disease-control cohort or lung-injury transport test qualified.','Age/sex support incomplete; observed male-only 3-versus-24 contrasts are sensitivity descriptions.','No meaningful-effect margin or confirmatory precision supplied.','No scientific acceptance or global RQ promotion.']})
    print(json.dumps({'expression_rows':len(expr),'composition_rows':len(comp),'microglia_cells':len(umap),'repertoire':rep,'source_overlap':overlap}))

if __name__=='__main__':main()
