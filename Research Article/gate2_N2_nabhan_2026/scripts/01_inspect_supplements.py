"""Inventory the seven cached source datasets without fitting or editing them."""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
from pathlib import Path
import zipfile


REPO = Path(__file__).resolve().parents[3]
EXPECTED_ARCHIVE_SHA256 = "d9d8b27e28d11b917f713598b009ac73bc6c393d24df231bd4bfe677a5876aee"


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--archive", type=Path, default=REPO / "tmp/rq-audit/sources/PMC13367804_supplements.zip")
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    if args.out.exists():
        parser.error(f"Refusing to overwrite {args.out}")
    from openpyxl import load_workbook  # Read-only schema inspection only.

    with args.archive.open("rb") as handle:
        archive_hash = hashlib.file_digest(handle, "sha256").hexdigest()
    if archive_hash != EXPECTED_ARCHIVE_SHA256:
        parser.error("Archive differs from P1 source recovery. Audit the new source before use.")
    destination = REPO / "raw_data/nabhan_2026_sources"
    destination.mkdir(parents=True, exist_ok=True)
    records = []
    with zipfile.ZipFile(args.archive) as archive:
        for index in range(1, 8):
            suffix = "csv" if index == 6 else "xlsx"
            name = f"pnas.2606113123.sd{index:02d}.{suffix}"
            content = archive.read(name)
            digest = hashlib.sha256(content).hexdigest()
            target = destination / name
            if target.exists():
                with target.open("rb") as handle:
                    if hashlib.file_digest(handle, "sha256").hexdigest() != digest:
                        raise ValueError(f"Existing source differs: {target.name}")
            else:
                target.write_bytes(content)
            item = {"dataset": f"S{index}", "path": target.relative_to(REPO).as_posix(),
                    "bytes": len(content), "sha256": digest}
            if suffix == "csv":
                rows = csv.reader(io.StringIO(content.decode("utf-8-sig")))
                header = next(rows)
                labels = [row[0] for row in rows]
                item.update(header=header, rows=len(labels), first_row_labels=labels[:5])
            else:
                workbook = load_workbook(target, read_only=True, data_only=True)
                item["sheets"] = workbook.sheetnames
                # Source dimensions are not evidence that all cells are populated.
                sheet = workbook.worksheets[0]
                item["first_sheet"] = {
                    "name": sheet.title, "declared_rows": sheet.max_row,
                    "declared_columns": sheet.max_column,
                    "header_preview": [list(row) for row in sheet.iter_rows(
                        min_row=1, max_row=1, max_col=12, values_only=True)],
                }
                workbook.close()
            records.append(item)
    report = {"schema": "nabhan-2026-supplement-schema/v1",
              "archive_sha256": archive_hash,
              "scope": "Source schema only; no biological endpoint computed. Headers may precede the true table header.",
              "datasets": records}
    args.out.parent.mkdir(parents=True, exist_ok=True)
    with args.out.open("x", encoding="utf-8") as handle:
        json.dump(report, handle, indent=2, ensure_ascii=False)
        handle.write("\n")
    print(json.dumps({"datasets": len(records), "out": str(args.out), "biological_models_run": 0}))


if __name__ == "__main__":
    main()
