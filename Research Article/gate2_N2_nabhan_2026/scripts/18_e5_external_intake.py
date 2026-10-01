"""Acquire GSE306184 and inspect design/schema without fitting expression effects."""
from pathlib import Path
import csv
import gzip
import hashlib
import io
import json
import tarfile
import urllib.request
from datetime import datetime, timezone
import openpyxl

HERE = Path(__file__).resolve().parents[1]
ROOT = HERE.parents[1]
CACHE = ROOT / "raw_data/GSE306184"
OUT = HERE / "runs/E5_metadata_v1"


def sha(data):
    return hashlib.sha256(data).hexdigest()


def main():
    if OUT.exists():
        raise SystemExit("Refusing to overwrite E5_metadata_v1")
    CACHE.mkdir(parents=True, exist_ok=True)
    sources = {
        "GSE306184_family.soft.gz": "https://ftp.ncbi.nlm.nih.gov/geo/series/GSE306nnn/GSE306184/soft/GSE306184_family.soft.gz",
        "GSE306184_RAW.tar": "https://ftp.ncbi.nlm.nih.gov/geo/series/GSE306nnn/GSE306184/suppl/GSE306184_RAW.tar",
    }
    inputs = []
    for name, url in sources.items():
        path = CACHE / name
        if not path.exists():
            request = urllib.request.Request(url, headers={"User-Agent": "Nb3-public-data-audit/1.0"})
            with urllib.request.urlopen(request, timeout=45) as response:
                data = response.read()
            path.write_bytes(data)
        inputs.append(dict(path=path.relative_to(ROOT).as_posix(), url=url,
                           bytes=path.stat().st_size, sha256=sha(path.read_bytes())))
    records, current = [], None
    for line in gzip.decompress((CACHE / "GSE306184_family.soft.gz").read_bytes()).decode().splitlines():
        if line.startswith("^SAMPLE = "):
            current = {"GSM": line.split(" = ", 1)[1]}
            records.append(current)
        elif current is not None and line.startswith("!Sample_") and " = " in line:
            key, value = line[8:].split(" = ", 1)
            current.setdefault(key, []).append(value)
    schemas = []
    with tarfile.open(CACHE / "GSE306184_RAW.tar") as archive:
        for member in archive.getmembers():
            if not member.isfile():
                continue
            data = archive.extractfile(member).read()
            if member.name.endswith(".gz"):
                data = gzip.decompress(data)
            if ".xlsx" not in member.name:
                raise ValueError(member.name)
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            sheets = []
            for ws in wb:
                rows = ws.iter_rows(values_only=True)
                sheets.append(dict(sheet=ws.title, rows=ws.max_row, columns=ws.max_column,
                                   header=list(next(rows))))
            schemas.append(dict(member=member.name, sha256_uncompressed=sha(data), sheets=sheets))
            wb.close()
    assert len(records) == 14
    OUT.mkdir(parents=True)
    (OUT / "sample_metadata.json").write_text(json.dumps(records, indent=2)+"\n", encoding="utf-8")
    (OUT / "workbook_schema.json").write_text(json.dumps(schemas, indent=2)+"\n", encoding="utf-8")
    fields = ["GSM", "title", "source_name_ch1", "characteristics_ch1", "treatment_protocol_ch1", "description", "data_processing", "supplementary_file"]
    with (OUT / "sample_manifest.tsv").open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, delimiter="\t")
        writer.writeheader()
        for record in records:
            writer.writerow({key: " | ".join(record.get(key, [])) if key != "GSM" else record[key] for key in fields})
    receipt = dict(accessed_utc=datetime.now(timezone.utc).isoformat(), inputs=inputs,
                   script_sha256=sha(Path(__file__).read_bytes()), sample_count=len(records),
                   workbook_count=len(schemas), status="metadata_and_schema_only_no_fit")
    (OUT / "run_record.json").write_text(json.dumps(receipt, indent=2)+"\n")
    print(json.dumps({"samples": len(records), "workbooks": len(schemas), "schemas": schemas[:1]}, indent=2))


if __name__ == "__main__":
    main()
