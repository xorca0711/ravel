"""Independent source checks using openpyxl, not the runner's XLSX XML parser."""
import csv
import hashlib
import json
import re
import xml.etree.ElementTree as ET
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[3]
CACHE = ROOT / 'raw_data/Yadav_followups_T0_20261031'
BR = ROOT / 'analysis/research/runs/Wg_P06_followup_t0_br_v1'
E = ROOT / 'analysis/research/runs/Yd1_P02_followup_t0_e_v1'
checks = []


def require(ok, name):
    if not ok:
        raise AssertionError(name)
    checks.append(name)


def read_csv(path):
    with path.open(encoding='utf-8', newline='') as stream:
        return list(csv.DictReader(stream))


raw = (CACHE / 'GSE273882_sample_metadata.txt').read_text(encoding='utf-8')
ids = re.findall(r'^\^SAMPLE = (GSM\d+)$', raw, re.M)
series = re.findall(r'^!Series_sample_id = (GSM\d+)$', (CACHE / 'GSE273882.txt').read_text(), re.M)
mapped = read_csv(BR / 'geo_sample_map.csv')
require(len(ids) == len(set(ids)) == 52 and set(ids) == set(series) == {r['gsm'] for r in mapped}, '52 exact unique GEO sample/series joins')
require(raw.count('!Sample_data_row_count = 0') == 52 and '!sample_table_begin' not in raw.lower(), 'GEO metadata contains no expression rows')
listing = json.loads((CACHE / 'hamanaka_figshare_listing.json').read_text())
bookpath = CACHE / 'hamanaka_metabolomics_v1.xlsx'
require(hashlib.md5(bookpath.read_bytes()).hexdigest() == listing['files'][0]['computed_md5'], 'Workbook matches source-provided MD5')
book = openpyxl.load_workbook(bookpath, read_only=False, data_only=False)
layout = json.loads((BR / 'workbook_layout.json').read_text(encoding='utf-8'))
require([x['sheet'] for x in layout] == book.sheetnames, 'All eleven source sheets inventoried in source order')
for item in layout:
    sheet = book[item['sheet']]
    labels = {c.coordinate: c.value for row in sheet for c in row if c.data_type == 's' and c.value is not None}
    require(labels == item['text_cells'], 'Independent string/coordinate agreement: ' + sheet.title)
samples = read_csv(BR / 'workbook_sample_map.csv')
source_samples = {(s.title, c.coordinate, c.value) for s in book for row in s for c in row
                  if c.column == 1 and isinstance(c.value, str) and re.search(r'_(?:0[1-4])$', c.value)}
require(source_samples == {(r['sheet'], r['source_cell'], r['sample_label']) for r in samples}, 'Exact workbook sample-label provenance, without treating repeated metabolites as units')
s = book['Figure 4F']
require(all(len({s.cell(r,1).value for r in range(start,start+6)}) == 6 for start in [3,11,19,27,35])
        and not any(isinstance(c.value,str) and c.value.endswith('_04') for row in s for c in row), '4F has three supplied labels per condition, no hidden fourth')
s = book['Figure 3F']
require(s['AP15'].value == s['AL15'].value == 'M+14' and s['AQ15'].value == s['AM15'].value == 'M+15', '3F duplicated isotope headers retained as issue')
s = book['Figure 4B']
require(s['P63'].value == s['A63'].value and s['T62'].value is None and s['T2'].value == 'M+4', '4B putrescine cannot inherit the earlier fraction header')
issues = read_csv(BR / 'source_issues.csv')
require({'H01','H02','H03','H04','H05','H06','H07','H08','L01','G01'} == {r['id'] for r in issues}, 'All source discrepancies and negatives retained')
tree = ET.parse(CACHE / 'atg5.xml').getroot()
fig6 = next(x for x in tree.findall('.//fig') if x.get('id') == 'fig6')
legend = ''.join(fig6.itertext())
require('(L and M)' in legend and 'n=5' in legend and 'n=3' in legend and 'n=4' in legend, 'Atg5 full legend includes distinct formation/RNA/staining units')
methods = ''.join(tree.itertext())
require('one sequencing library per group' in methods and 'On Day 10' in methods, 'Atg5 pooling and timing limits verified in primary text')
endpoints = read_csv(E / 'endpoint_unit_map.csv')
require({r['source'] for r in endpoints} == {'Atg5_2026','A13','A20','A8','A19','A22','Yadav2025'}, 'Separate owner endpoint map preserves every required repository source')
require(any('PDPN+/SPC+' in r['denominator'] and 'absolute' in r['interpretation_limit'] for r in endpoints), 'Exact staining denominator and absolute-output limit retained')
for run in [BR,E]:
    access = json.loads((run / 'source_access.json').read_text())
    require(all(not r['content_verified'] for r in access['requests'] if r['file'].endswith('_attempt.pdf')), 'Challenge HTML rejected as PDF: ' + run.name)
print(json.dumps({'ok': True, 'checks': len(checks), 'methods': ['Independent openpyxl/source XML/GEO reread', 'Primary figure and supplement visual review is recorded in T0_RESULTS.md'], 'labels': checks}, indent=2))
