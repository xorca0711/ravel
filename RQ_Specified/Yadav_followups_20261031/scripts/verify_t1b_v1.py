"""New-endpoint boundary tests and independent source/arithmetic verification."""
import argparse
import csv
import hashlib
import json
import math
from collections import defaultdict
from pathlib import Path

ROOT=Path(__file__).resolve().parents[3]
HERE=Path(__file__).resolve().parents[1]
BOOK=ROOT/'raw_data/Yadav_followups_T0_20261031/hamanaka_metabolomics_v1.xlsx'
checks=0


def require(ok,message):
    global checks
    if not ok:raise AssertionError(message)
    checks+=1


def read(path):
    with path.open(encoding='utf-8',newline='') as f:return list(csv.DictReader(f))


def close(a,b,scale=None):
    a,b=float(a),float(b)
    scale=max(abs(a),abs(b),1e-300) if scale is None else max(abs(scale),1e-300)
    return abs(a-b)<=64*math.ulp(scale)


def boundaries():
    from describe_t1b_v1 import PERCENT,AREA,parse_channel,channel_summary,vector_diagnostic,header_issues
    def c(raw):return {'raw':raw,'text':raw,'kind':'n','formula':''}
    def rec(raw,i):return {**parse_channel(c(raw),PERCENT),'sample_label':str(i),'header_status':'qualified'}
    rr=[rec('25',i) for i in range(4)]
    require(channel_summary(rr,4)['mean']=='25','full four-label channel qualifies')
    v=vector_diagnostic([c('25'),c('')])
    require(v['completeness'].startswith('incomplete') and v['complete_sum_reported_percent']=='' and v['partial_valid_channel_subtotal_percent']=='25','missing neighbor is incomplete, not zero')
    require(channel_summary(rr,4)['status']=='qualified','neighbor does not gate fully supplied target channel')
    rr[3]=rec('',3);s=channel_summary(rr,4)
    require(s['status']=='held' and s['mean']=='' and s['n_supplied']==4 and s['n_valid']==3,'missing target retains n and holds mean')
    require(channel_summary(rr[:3],4)['status']=='held','partial supplied n cannot qualify')
    require(channel_summary(rr[:3],3)['status']=='qualified','three-label 4F has explicit n=3')
    require(parse_channel(c('100.01'),PERCENT)['value_status']=='outside_declared_percent_domain','above 100 held without clipping')
    require(parse_channel(c('100.01'),AREA)['value_status']=='valid','area does not inherit percent domain')
    require(parse_channel(c('.5'),PERCENT)['numeric_value']=='0.5','literal percent not rescaled from outcome')
    require(vector_diagnostic([c('50'),c('50.000000001')])['closure_diagnostic'].endswith('does_not_close_at_fixed_tolerance'),'failed old precision rule not relaxed')
    require(channel_summary([rec('50.000000001',i) for i in range(4)],4)['status']=='qualified','valid channel remains described despite failed vector closure')
    bad=rec('25',3);bad['header_status']='held'
    require(channel_summary([rec('25',i) for i in range(3)]+[bad],4)['status']=='held','unit/header mismatch holds summary')
    require(vector_diagnostic([c('0'),c('0')])['complete_sum_reported_percent']=='0','complete zero vector retained without normalization')
    require(parse_channel(c('-1'),PERCENT)['value_status']=='negative','negative percent retained invalid')
    badc=c('20');badc['formula']='10+10'
    require(parse_channel(badc,PERCENT)['value_status']=='formula_cached_value_not_used','formula is not an observed channel')
    print(json.dumps({'ok':True,'new_endpoint_boundary_checks':checks,'prior_parser_boundary_checks':'reused, unchanged'}))


def verification():
    import openpyxl
    from openpyxl.utils import column_index_from_string as ci,get_column_letter as cl
    book=openpyxl.load_workbook(BOOK,data_only=False)
    oldpath=ROOT/'analysis/research/runs/Wg_P06_followup_t1a_r1_v1/observation_qc.csv'
    oldqc={(r['block'],r['source_row']):r for r in read(oldpath)}
    results={}
    for job in ['r1','r2']:
        spec=json.loads((HERE/f'config/t1b_{job}_spec_v1.json').read_text());run=ROOT/f'analysis/research/runs/Wg_P06_followup_t1b_{job}_v1'
        rows=read(run/'source_values.csv');ss=read(run/'summaries.csv');qq=read(run/'vector_qc.csv');units=read(run/'unit_and_header_map.csv')
        expected={};bmap={b['id']:b for b in spec['blocks']}
        header_ok={}
        for b in spec['blocks']:
            sh=book[b['sheet']];okay=True
            for m in b['measures']:
                okay &= sh[m['unit_cell']].value==m['unit_literal']
                for address,label in m['isotope_headers'].items():okay &= sh[address].value==label
            for obs in b['observations']:
                okay &= sh['A'+str(obs['row'])].value==obs['sample_label']
                okay &= sh[b['metabolite_column']+str(obs['row'])].value==b['metabolite']
                for m in b['measures']:
                    for n in range(ci(m['columns'][0]),ci(m['columns'][1])+1):
                        addr=cl(n)+str(obs['row']);expected[(b['sheet'],addr)]=(b,m,obs)
            header_ok[b['id']]=okay
        require(len(rows)==len(expected) and {(r['sheet'],r['source_cell']) for r in rows}==set(expected),job+' exact authorized cell set')
        require(len(units)==sum(len(b['measures']) for b in spec['blocks']),job+' all measure maps')
        for u in units:
            require((u['header_status']=='qualified')==header_ok[u['block']],job+' independently validated map')
            require(book[u['sheet']][u['unit_cell']].value==u['observed_unit_literal']==u['expected_unit_literal'],job+' explicit unit provenance')
        independent={}
        for r in rows:
            b,m,obs=expected[(r['sheet'],r['source_cell'])];sh=book[r['sheet']];v=sh[r['source_cell']].value
            status=('missing' if v is None else 'formula_cached_value_not_used' if isinstance(v,str) and v.startswith('=') else
                    'nonnumeric' if isinstance(v,bool) or not isinstance(v,(int,float)) else 'nonfinite' if not math.isfinite(v) else
                    'negative' if v<0 else 'outside_declared_percent_domain' if m['name']=='reported_channel_percent' and v>100 else 'valid')
            require(status==r['value_status'],job+' independent cell domain '+r['source_cell'])
            require(r['measure']==m['name'] and r['reported_condition']==obs['condition'] and r['sample_label']==obs['sample_label'],job+' exact cell/condition/label map')
            require(sh[r['unit_cell']].value==r['source_unit_literal']==m['unit_literal'],job+' literal unit')
            require(sh[r['header_cell']].value==r['observed_isotope_header']==r['isotope'],job+' literal isotope')
            if status in ['valid','negative','outside_declared_percent_domain']:require(close(v,r['numeric_value']),job+' independent numeric read')
            key=(r['block'],r['measure'],r['isotope'],r['reported_condition'])
            independent.setdefault(key,[]).append((v,status,r))
        index={};means={}
        for s in ss:
            key=(s['block'],s['measure'],s['isotope'],s['reported_condition']);rr=independent[key];n=bmap[s['block']]['n_per_condition']
            valid=[v for v,status,r in rr if status=='valid'];okay=len(rr)==len(valid)==n and header_ok[s['block']]
            require(int(s['n_supplied'])==len(rr)==n and int(s['n_valid'])==len(valid) and int(s['n_expected'])==n,job+' full supplied n')
            require((s['status']=='qualified')==okay,job+' channel eligibility independent of vector')
            if okay:
                mean=math.fsum(valid)/n;means[key]=mean
                require(close(mean,s['mean']) and close(min(valid),s['minimum']) and close(max(valid),s['maximum']),job+' independent channel arithmetic')
            else:require(s['mean']==s['minimum']==s['maximum']=='',job+' held target not partial-n summarized')
            index[key]=s
        require(len(index)==len(independent)==len(ss),job+' all channel summaries retained')
        for d in read(run/'descriptive_differences.csv'):
            k=(d['block'],d['measure'],d['isotope']);a,b=index[k+('UT',)],index[k+('TGF-beta',)];okay=a['status']==b['status']=='qualified'
            require((d['status']=='qualified')==okay,job+' difference eligibility')
            if okay:
                av,bv=means[k+('UT',)],means[k+('TGF-beta',)]
                require(close(bv-av,d['difference_native_units'],max(abs(av),abs(bv))),job+' independent difference arithmetic')
            else:require(d['difference_native_units']=='',job+' held difference')
        for q in qq:
            b=bmap[q['block']];m=next(m for m in b['measures'] if m['name']=='reported_channel_percent');sh=book[b['sheet']];row=int(q['source_row'])
            vs=[sh.cell(row,n).value for n in range(ci(m['columns'][0]),ci(m['columns'][1])+1)]
            valid=[v for v in vs if isinstance(v,(float,int)) and not isinstance(v,bool) and math.isfinite(v) and 0<=v<=100]
            complete=len(valid)==len(vs);tol=64*len(vs)*math.ulp(100.0)
            require(int(q['n_valid_percent_channels'])==len(valid) and int(q['n_channels_expected'])==len(vs),job+' vector completeness count')
            require(q['completeness'].startswith('complete_')==complete,job+' incomplete not called complete')
            if complete:
                total=math.fsum(valid);status='complete:closes_at_fixed_tolerance' if abs(total-100)<=tol else 'complete:does_not_close_at_fixed_tolerance'
                require(close(total,q['complete_sum_reported_percent']) and close(total-100,q['residual_from_100_percentage_points'],100),job+' complete sum/residual')
                require(q['partial_valid_channel_subtotal_percent']=='' and q['closure_diagnostic']==status,job+' unchanged fixed precision rule')
            else:
                require(q['complete_sum_reported_percent']==q['residual_from_100_percentage_points']=='' and q['closure_diagnostic']=='incomplete:not_tested',job+' no partial sum called complete')
                require(close(math.fsum(valid),q['partial_valid_channel_subtotal_percent']) if valid else q['partial_valid_channel_subtotal_percent']=='',job+' explicitly incomplete subtotal')
            require(close(tol,q['fixed_tolerance_at_100']),job+' tolerance frozen independently of values')
            if job=='r1':require(json.loads(q['unchanged_v1_diagnostic_json'])==oldqc[(q['block'],q['source_row'])],'every original v1 diagnostic unchanged')
            else:require(q['unchanged_v1_diagnostic_json']=='','R2 has no fabricated historical vector check')
        expected_vectors=sum(len(b['observations']) for b in spec['blocks'] if any(m['name']=='reported_channel_percent' for m in b['measures']))
        require(len(qq)==expected_vectors and len({(q['block'],q['source_row']) for q in qq})==expected_vectors,job+' complete vector QC coverage')
        disp=json.loads((run/'disposition.json').read_text());require(hashlib.sha256(BOOK.read_bytes()).hexdigest()==disp['source_sha256'],job+' source hash')
        require(disp['literal_cells']==len(rows) and disp['summaries']==len(ss) and disp['held_summaries']==sum(s['status']=='held' for s in ss),job+' disposition matches tables')
        if job=='r1':require(hashlib.sha256(oldpath.read_bytes()).hexdigest()==disp['old_v1_qc_sha256'],'unchanged original QC file hash')
        results[job]={'source_cells':len(rows),'summaries':len(ss),'held_summaries':sum(s['status']=='held' for s in ss),'vector_observations':len(qq)}
    print(json.dumps({'ok':True,'independent_checks':checks,'results':results,
                      'method':'Independent openpyxl source/unit/header reading, math.fsum arithmetic, percent-domain and full-n checks; unchanged v1 diagnostic equality'},indent=2))


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--self-test',action='store_true');args=parser.parse_args()
    boundaries() if args.self_test else verification()
