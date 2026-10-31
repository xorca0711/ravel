"""Boundary tests and independent openpyxl/float arithmetic verification."""
import argparse
import csv
import hashlib
import json
import math
from collections import defaultdict
from decimal import Decimal
from pathlib import Path

ROOT=Path(__file__).resolve().parents[3]
HERE=Path(__file__).resolve().parents[1]
BOOK=ROOT/'raw_data/Yadav_followups_T0_20261031/hamanaka_metabolomics_v1.xlsx'
checks=0


def require(ok,message):
    global checks
    if not ok:raise AssertionError(message)
    checks+=1


def boundary_tests():
    from t1a_common_v1 import fraction_candidates, literal_fields, summary
    def c(x):return {'raw':x,'text':x,'kind':'n','formula':''}
    require(fraction_candidates([c('.4'),c('.6')])['candidates']==['1'],'unit fraction recognized')
    require(fraction_candidates([c('40'),c('60')])['candidates']==['100'],'percent recognized')
    require(fraction_candidates([c('.4'),c('.61')])['candidates']==[],'non-closing row held without renormalization')
    require(fraction_candidates([c('0'),c('0')])['candidates']==[],'zero denominator not called 100 percent')
    require(fraction_candidates([c('-0.01'),c('1.01')])['status']=='invalid_fraction_cell','negative fraction cannot pass closure')
    require(literal_fields(c('-1'))['numeric_value']=='-1','negative source value preserved')
    require(literal_fields(c('NaN'))['value_status']=='nonfinite','nonfinite source value held')
    require(literal_fields({'raw':'1','text':'1','kind':'n','formula':'1+0'})['value_status']=='formula_cached_value_not_used','cached formula not silently used')
    good=literal_fields(c('2'));bad=literal_fields(c('-1'))
    s=summary([good,bad]);require(s['n_supplied']==2 and s['n_valid']==1 and s['mean']=='' and s['status']=='held','invalid observation cannot disappear into a reduced-n mean')
    s=summary([good,literal_fields(c('4'))]);require(s['mean']=='3' and s['minimum']=='2' and s['maximum']=='4','mean and observed range')
    print(json.dumps({'ok':True,'boundary_checks':checks}))


def csv_rows(path):
    with path.open(encoding='utf-8',newline='') as f:return list(csv.DictReader(f))


def near(a,b):
    a,b=float(a),float(b)
    return abs(a-b)<=64*math.ulp(max(abs(a),abs(b),1e-300))


def run_verification():
    import openpyxl
    book=openpyxl.load_workbook(BOOK,read_only=False,data_only=False)
    summaries_all={}
    for job in ['r1','b1']:
        run=ROOT/f'analysis/research/runs/Wg_P06_followup_t1a_{job}_v1'
        spec=json.loads((HERE/f'config/t1a_{job}_spec_v1.json').read_text())
        rows=csv_rows(run/'source_values.csv');summaries=csv_rows(run/'summaries.csv')
        expected=set()
        if job=='r1':
            from openpyxl.utils import column_index_from_string as ci,get_column_letter as cl
            for b in spec['blocks']:
                for start,end in [b['area_columns'],b['fraction_columns']]:
                    for col in range(ci(start),ci(end)+1):
                        for row in b['rows']:expected.add((b['sheet'],f'{cl(col)}{row}'))
        else:
            for p in spec['panels']:
                for col in p['sample_headers']:
                    for row in p['rows']:expected.add((p['sheet'],col+row))
        require(expected=={(r['sheet'],r['source_cell']) for r in rows} and len(rows)==len(expected),job+' exact authorized source-cell set')
        for r in rows:
            v=book[r['sheet']][r['source_cell']].value
            status='missing' if v is None else 'formula_cached_value_not_used' if isinstance(v,str) and v.startswith('=') else 'nonnumeric' if not isinstance(v,(int,float)) else 'nonfinite' if not math.isfinite(v) else 'negative' if v<0 else 'valid'
            require(status==r['value_status'],job+' independent cell status '+r['source_cell'])
            if status in ['valid','negative']:
                require(near(v,r['numeric_value']),job+' independent literal numeric read '+r['source_cell'])
            require(str(book[r['sheet']][r['header_cell']].value)==(r['source_header'] if job=='r1' else r['full_sample_label']),job+' exact header '+r['header_cell'])
        grouped=defaultdict(list)
        for r in rows:
            key=(r['block'],r['measure'],r['isotope'],r['reported_condition']) if job=='r1' else (r['sheet'],r['metabolite'],r['native_group'])
            grouped[key].append(r)
        index={}
        for s in summaries:
            key=(s['block'],s['measure'],s['isotope'],s['reported_condition']) if job=='r1' else (s['sheet'],s['metabolite'],s['native_group'])
            rr=grouped[key];index[key]=s
            valid=[r for r in rr if r['value_status']=='valid' and (job!='r1' or r['measure']!='reported_fraction' or r['fraction_status']=='qualified')]
            require(int(s['n_supplied'])==len(rr)==4 and int(s['n_valid'])==len(valid),job+' source count '+str(key))
            if s['status']=='qualified':
                require(len(valid)==len(rr),job+' no selective deletion '+str(key))
                values=[float(book[r['sheet']][r['source_cell']].value) for r in rr]
                require(near(math.fsum(values)/len(values),s['mean']) and near(min(values),s['minimum']) and near(max(values),s['maximum']),job+' independent arithmetic '+str(key))
            else:require(s['mean']==s['minimum']==s['maximum']=='',job+' held summary left blank '+str(key))
        if job=='r1':
            qc=csv_rows(run/'observation_qc.csv');units=csv_rows(run/'block_units.csv')
            source_candidates={}
            for b in spec['blocks']:
                candidates={}
                for row in b['rows']:
                    vs=[book[b['sheet']].cell(row,col).value for col in range(ci(b['fraction_columns'][0]),ci(b['fraction_columns'][1])+1)]
                    okay=all(isinstance(v,(int,float)) and math.isfinite(v) and v>=0 for v in vs)
                    cand=[str(scale) for scale in [1,100] if okay and abs(math.fsum(vs)-scale)<=64*len(vs)*math.ulp(float(scale))]
                    candidates[row]=cand
                    q=next(q for q in qc if q['block']==b['id'] and int(q['source_row'])==row)
                    require(q['scale_candidates']=='|'.join(cand),'independent fraction scale/closure '+b['id']+str(row))
                single={v[0] for v in candidates.values() if len(v)==1};scale=next(iter(single)) if len(single)==1 else ''
                unit=next(u for u in units if u['block']==b['id'])
                require(unit['fraction_scale']==scale,'independent block scale '+b['id'])
                for row,cand in candidates.items():
                    q=next(q for q in qc if q['block']==b['id'] and int(q['source_row'])==row)
                    expected_status=bool(scale and cand==[scale] and unit['header_status']=='qualified' and q['reported_condition']!='unknown')
                    require((q['fraction_status']=='qualified')==expected_status,'independent held row '+b['id']+str(row))
            for d in csv_rows(run/'descriptive_differences.csv'):
                key=(d['block'],d['measure'],d['isotope'])
                a,b=index[key+('UT',)],index[key+('TGF-beta',)]
                okay=a['status']==b['status']=='qualified'
                require((d['status']=='qualified')==okay,'descriptive contrast eligibility '+str(key))
                if okay:require(near(Decimal(b['mean'])-Decimal(a['mean']),d['difference_native_units']),'descriptive contrast arithmetic '+str(key))
                else:require(d['difference_native_units']=='','held contrast blank')
        else:
            require(len(rows)==144 and len(summaries)==36,'all B1 native groups retained')
            require(all(r['condition_key'].startswith('unresolved') for r in rows),'no prefix interpreted as medium')
        disposition=json.loads((run/'disposition.json').read_text())
        require(disposition['source_sha256']==hashlib.sha256(BOOK.read_bytes()).hexdigest(),'source hash '+job)
        summaries_all[job]={'source_values':len(rows),'summaries':len(summaries),'held_summaries':sum(r['status']=='held' for r in summaries)}
    print(json.dumps({'ok':True,'independent_checks':checks,'results':summaries_all,'method':'Independent openpyxl source reading and math.fsum/min/max; separate binary64 closure and source-cell membership checks'},indent=2))


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--self-test',action='store_true');a=p.parse_args()
    boundary_tests() if a.self_test else run_verification()
