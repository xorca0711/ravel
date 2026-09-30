"""Independent source-coordinate, arithmetic and figure-integrity checks for A19.

Recomputes from workbook cells and downloaded count columns; does not import
analysis functions. Requires the ignored cache produced by scripts 00-02.
"""
from pathlib import Path
import hashlib, json, platform
from datetime import datetime, timezone
import numpy as np
import pandas as pd
import openpyxl
from scipy.stats import t
import pymupdf
B = Path(__file__).resolve().parents[1]
T = B / 'tables/exploratory_v1'
F = B / 'figures/exploratory_v1'
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def read(name): return pd.read_csv(T / name, sep='\t')
def close(a, b): np.testing.assert_allclose(a, b, rtol=1e-7, atol=1e-7, equal_nan=True)
def main():
    config = json.loads((B/'config/exploratory_v1.json').read_text())
    frozen = json.loads((B/'reports/exploratory_v1_contract.json').read_text())
    assert frozen['contract_sha256'] == sha(B/'config/exploratory_v1.json')
    manifest = json.loads((B/'metadata/download_manifest.json').read_text())
    for r in manifest:
        assert sha(B/'cache'/r['file']) == r['sha256'] == frozen['input_sha256'][r['file']]
    raw = read('qpcr_source_rows.tsv')
    w = openpyxl.load_workbook(B/'cache/adult_source3.xlsx', data_only=True)
    source_sheets = {x.strip(): x for x in config['qpcr']['sheets']}
    source_values = {}
    for row in raw.itertuples():
        ws = w[source_sheets[row.sheet]]
        assert ws.cell(row.source_row, 3).value == row.donor
        for column, expected in [(1, row.source_gene), (2, row.condition)]:
            i = row.source_row
            while ws.cell(i, column).value is None: i -= 1
            assert str(ws.cell(i, column).value).strip() == expected
        x = ws.cell(row.source_row, 4).value
        if isinstance(x, (int, float)):
            close(x, row.delta_ct)
        else:
            assert pd.isna(row.delta_ct) and str(x).startswith('undet')
            x = np.nan
        source_values[(row.sheet, row.gene, row.condition, row.donor)] = x
    assert len(raw) == 143 and (raw.ct_status=='undetected').sum() == 3
    effects = read('qpcr_paired_effects.tsv')
    audit = read('qpcr_pairing_audit.tsv')
    for r in audit.itertuples():
        a = (r.sheet, r.gene, r.comparison, r.donor)
        z = (r.sheet, r.gene, r.reference, r.donor)
        expected = ('missing_arm' if a not in source_values or z not in source_values
                    else 'nonquantified_Ct' if not np.isfinite(source_values[a]+source_values[z])
                    else 'complete')
        assert r.status == expected
    assert len(effects) == (audit.status=='complete').sum() == 72
    for r in effects.itertuples():
        value = source_values[(r.sheet,r.gene,r.reference,r.donor)] - source_values[(r.sheet,r.gene,r.comparison,r.donor)]
        close(r.effect_log2, value)
    keys = ['sheet','gene','comparison','reference']
    for r in read('qpcr_effect_summary.tsv').itertuples():
        mask = np.logical_and.reduce([effects[k]==getattr(r,k) for k in keys])
        v = effects.loc[mask,'effect_log2'].to_numpy()
        assert r.n_pairs == len(v)
        if len(v):
            close([r.mean_log2,r.median_log2,r.min_log2,r.max_log2], [v.mean(),np.median(v),v.min(),v.max()])
        if len(v) >= 2:
            half = t.ppf(.975, len(v)-1)*v.std(ddof=1)/np.sqrt(len(v))
            close([r.ci95_low,r.ci95_high],[v.mean()-half,v.mean()+half])
        else: assert pd.isna(r.ci95_low) and pd.isna(r.ci95_high)
    w.close()
    meta = read('bulk_samples.tsv')
    assert len(meta)==8 and set(meta.gsm)=={f'GSM{x}' for x in range(5934407,5934415)}
    count = pd.read_csv(B/'cache/bulk_counts.tsv.gz',sep='\t',index_col=0)
    for r in meta.itertuples():
        source = pd.read_csv(B/'cache'/r.url.split('/')[-1],sep='\t',comment='#',usecols=[0,6],index_col=0)
        pd.testing.assert_index_equal(count.index, source.index, check_names=False)
        close(count[r.title].to_numpy(),source.iloc[:,0].to_numpy())
    assert count.shape == (60662,8)
    keep = (count>=10).sum(axis=1)>=2
    assert keep.sum()==14581
    filtered=count.loc[keep]
    norms=read('bulk_normalization.tsv').set_index('sample').loc[count.columns]
    close(norms.library_size,filtered.sum())
    close(norms.library_size*norms.tmm_factor,norms.effective_library_size)
    close(np.exp(np.log(norms.tmm_factor).mean()),1)
    lookup=json.loads((B/'metadata/human_marker_lookup.json').read_text())['records']
    members=read('bulk_panel_membership.tsv')
    expr=read('bulk_marker_expression.tsv')
    scores=read('bulk_program_scores.tsv')
    computed={}
    for method, filename in [('TMM','bulk_tmm_log2cpm.tsv.gz'),('total_count','bulk_total_log2cpm.tsv.gz')]:
        denom = norms.effective_library_size if method=='TMM' else norms.library_size
        expected=np.log2(filtered.div(denom,axis=1)*1e6+1)
        written=pd.read_csv(B/'cache'/filename,sep='\t',index_col=0).loc[expected.index,expected.columns]
        close(expected.values,written.values)
        expected.index=expected.index.str.replace(r'\.\d+$','',regex=True)
        computed[method]=expected
        for r in expr[expr.method==method].itertuples(): close(r.log2cpm,expected.loc[lookup[r.gene]['id'],r.sample])
        for panel,genes in config['bulk']['panels'].items():
            ids=[lookup[g]['id'] for g in genes if lookup[g]['id'] in expected.index]
            eligible=len(ids)>=2 and len(ids)/len(genes)>=.6
            z=scores[(scores.method==method)&(scores.panel==panel)]
            assert len(z)==(8 if eligible else 0)
            if eligible:
                for r in z.itertuples():close(r.score,expected.loc[ids,r.sample].mean())
            if method=='TMM':
                actual=members[members.panel==panel].set_index('gene')
                for g in genes: assert bool(actual.loc[g,'in_filtered_matrix'])==(lookup[g]['id'] in expected.index)
    program=read('bulk_program_effects.tsv')
    gene_effects=read('bulk_marker_effects.tsv')
    for frame,source,variable,value in [(program,scores,'panel','score'),(gene_effects,expr,'gene','log2cpm')]:
        for r in frame.itertuples():
            subset=source[(source.method==r.method)&(source[variable]==getattr(r,variable))&(source.block==r.block)&(source.background==r.background)].set_index('input')
            close(r.absence_effect,subset.loc['no_CHIR',value]-subset.loc['CHIR',value])
    for r in read('bulk_response_interactions.tsv').itertuples():
        z=program[(program.method==r.method)&(program.panel==r.panel)&(program.block==r.block)].set_index('background')
        close(r.GSK3KD_minus_control_response,z.loc['GSK3KD','absence_effect']-z.loc['control','absence_effect'])
    wide=program.pivot(index=['panel','block','background'],columns='method',values='absence_effect').reset_index()
    flips=wide[np.sign(wide.TMM)!=np.sign(wide.total_count)]
    render=json.loads((B/'reports/figure_render.json').read_text())
    assert render['script_sha256']==sha(B/'scripts/04_figures.py')
    for name,digest in render['input_tables'].items(): assert sha(T/name)==digest
    assert len(render['outputs'])==12
    for name,digest in render['outputs'].items(): assert sha(F/name)==digest
    for p in F.glob('*.pdf'):
        d=pymupdf.open(p);assert len(d)==1 and len(d[0].get_text())>200;d.close()
    report={'status':'passed','verified_utc':datetime.now(timezone.utc).isoformat(),
            'checks':['frozen contract and raw source hashes','143 source cell coordinates and donor keys','72 complete paired contrasts and intervals','8 source count columns and 14581-gene filter','TMM denominator arithmetic and full normalized matrices','marker coverage, program and gene effects, interactions','12 export hashes and 4 single-page text-bearing PDFs'],
            'limits':'Checks do not establish source biological independence, PCR efficiency, TMM assumptions or mature fate/function.',
            'pairing_status':audit.status.value_counts().to_dict(),
            'normalization_sign_disagreements':flips.to_dict(orient='records'),
            'versions':{'python':platform.python_version(),'numpy':np.__version__,'pandas':pd.__version__,'openpyxl':openpyxl.__version__},
            'script_sha256':{p.name:sha(p) for p in sorted((B/'scripts').glob('*')) if p.is_file()},
            'table_sha256':{p.name:sha(p) for p in sorted(T.glob('*.tsv'))}}
    (B/'reports/verification.json').write_text(json.dumps(report,indent=2)+'\n',encoding='utf-8')
    print(f'PASS: source, arithmetic and export integrity; {len(flips)} of 20 program effects change sign under total-count sensitivity.')
    print(flips.to_string(index=False))
if __name__=='__main__':main()
