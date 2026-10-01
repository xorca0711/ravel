"""A23 exploratory P2/P3 analysis. All outcomes are descriptive and source-specific."""
import argparse
import collections
import csv
import datetime
import hashlib
import json
import platform
from pathlib import Path

ROOT = Path(__file__).resolve().parents[3]
RQ = ROOT / 'RQ_Specified/A23_slc34a2_transition_homeostasis'
RUN = 'external_pilot_v2'

def digest(p):
    h = hashlib.sha256()
    with p.open('rb') as f:
        for b in iter(lambda: f.read(1048576), b''): h.update(b)
    return h.hexdigest()

def read(p): return json.loads(p.read_text(encoding='utf-8'))

def write_table(path, rows):
    if not rows: raise ValueError('Empty table: '+str(path))
    with path.open('w', encoding='utf-8', newline='') as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]), delimiter='\t', lineterminator='\n')
        w.writeheader(); w.writerows(rows)

def annotation_key(library, barcode):
    # Only the observed 10x gem suffix -1 is removable; preserve other identifiers.
    if not barcode.endswith('-1'): raise ValueError('Unexpected gem suffix')
    return library, barcode[:-2]

def annotation_map(rows, prefixes):
    result = {}
    for row in rows:
        prefix, barcode = row[0].split('.', 1)
        key = (prefixes[prefix], barcode)
        if key in result: raise ValueError('Duplicate annotation key')
        result[key] = row[1]
    return result

def check(tracked_only=False):
    receipt = read(RQ/'metadata'/RUN/'run_record.json')
    for entry in receipt['inputs']+receipt['outputs']:
        if tracked_only and entry['path'].startswith('raw_data/'): continue
        p = ROOT/entry['path']
        if digest(p) != entry['sha256']: raise ValueError('Hash mismatch: '+entry['path'])
    print('A23 external pilot receipt PASS'+(' (tracked files only)' if tracked_only else ''))

def run():
    import h5py
    import numpy as np
    import openpyxl
    import scipy
    from scipy.sparse import csc_matrix
    paths = {k: RQ/k/RUN for k in ['tables','metadata']}
    for p in paths.values():
        if p.exists(): raise FileExistsError(p)
    cfg = read(RQ/'config/external_pilot_v2.json')
    add = read(RQ/'config/external_pilot_v1_annotation_addendum.json')
    bio = read(RQ/'config/biochemical_audit_v1.json')
    acq = RQ/'metadata/external_acquisition_v1'
    matrices = read(acq/'matrices.json')
    sources = read(acq/'sources.json')
    for r in sources['source_files']+matrices:
        if digest(ROOT/r['path']) != r['sha256']: raise ValueError('Changed source: '+r['path'])
    workbook = ROOT/next(r['path'] for r in sources['source_files'] if r['path'].endswith('.xlsx'))
    book = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    annotations = annotation_map(list(book['Fig 2a'].values)[1:], add['prefix_map'])
    selected_map = read(acq/'selected_gene_map.json')['mappings']
    for p in paths.values(): p.mkdir(parents=True)
    tables = paths['tables']
    coverage=[]; summaries=[]; genes=[]; cross=[]; per_cell=[]; annotation_coverage=[]
    readouts = [(role,g) for role,gs in cfg['readouts'].items() for g in gs]
    selection = cfg['epithelial_features']+cfg['at2_selection_features']+sum(cfg['competing_lineages'].values(),[])
    allgenes = list(dict.fromkeys(selection+[g for _,g in readouts]))
    assert not set(selection).intersection(g for _,g in readouts)
    for rec in matrices:
        library=rec['library']
        with h5py.File(ROOT/rec['path'],'r') as f:
            g=f['hg19']; symbols=np.char.decode(g['gene_names'][:]); ids=np.char.decode(g['genes'][:]); barcodes=np.char.decode(g['barcodes'][:])
            x=csc_matrix((g['data'][:],g['indices'][:],g['indptr'][:]),shape=tuple(g['shape'][:])).tocsr()
        x.sum_duplicates(); x.eliminate_zeros()
        assert np.all(x.data>=0) and np.all(x.data==np.floor(x.data))
        assert len(set(barcodes))==len(barcodes)
        total=np.asarray(x.sum(axis=0)).ravel(); ngenes=np.asarray((x>0).sum(axis=0)).ravel()
        mito=np.asarray(x[np.char.startswith(symbols,'MT-')].sum(axis=0)).ravel()/np.maximum(total,1)
        vectors={}
        for gene in allgenes:
            indices=np.flatnonzero(symbols==gene)
            genes.append(dict(library=library,gene=gene,n_features=len(indices),stable_ids=';'.join(ids[indices])))
            if len(indices)>1: raise ValueError('Ambiguous symbol: '+gene)
            vectors[gene]=x[indices[0]].toarray().ravel() if len(indices) else None
            if gene in selection and len(indices)!=1: raise ValueError('Missing annotation feature '+gene)
        for m in selected_map:
            ii=np.flatnonzero(symbols==m['human'])
            if len(ii)!=1 or ids[ii[0]] not in m['ensembl_ids']: raise ValueError('Readout stable-ID mismatch')
        def detected(features): return sum(vectors[g]>=cfg['detection_min_umi'] for g in features)
        epi=detected(cfg['epithelial_features'])>=cfg['epithelial_min_features']
        at2=detected(cfg['at2_selection_features'])>=cfg['at2_min_features']
        competing={k:detected(v)>=cfg['competing_min_features'][k] for k,v in cfg['competing_lineages'].items()}
        mixed=np.logical_or.reduce(list(competing.values()))
        labels=np.array([annotations.get(annotation_key(library,b),'unmatched') for b in barcodes])
        present_keys={annotation_key(library,b) for b in barcodes}
        source_keys={k for k in annotations if k[0]==library}
        annotation_coverage.append(dict(library=library,filtered_barcodes=len(barcodes),published_barcodes=len(source_keys),matched=len(source_keys & present_keys),published_not_filtered=len(source_keys-present_keys),filtered_not_published=len(present_keys-source_keys)))
        for qc_name,q in cfg['qc'].items():
            qc=(ngenes>=q['min_genes'])&(ngenes<=q['max_genes'])&(total<=q['max_counts'])&(mito<=q['max_mito_fraction'])
            eligible_epi=qc & epi & ~mixed
            groups={'epithelial_candidate':eligible_epi,'AT2_candidate':eligible_epi & at2,'AT2_published_intersection':eligible_epi & at2 & (labels=='AT2')}
            coverage.append(dict(library=library,qc=qc_name,filtered_cells=len(barcodes),qc_cells=int(qc.sum()),epithelial_pre_exclusion=int((qc & epi).sum()),colineage_excluded=int((qc & epi & mixed).sum()),epithelial_candidate=int(eligible_epi.sum()),AT2_candidate=int(groups['AT2_candidate'].sum()),AT2_published_intersection=int(groups['AT2_published_intersection'].sum()),median_umi_qc=float(np.median(total[qc])),median_genes_qc=float(np.median(ngenes[qc]))))
            for group,mask in groups.items():
                n=int(mask.sum())
                for label,count in collections.Counter(labels[mask]).items(): cross.append(dict(library=library,qc=qc_name,group=group,published_label=label,n=count))
                for role,gene in readouts:
                    v=vectors[gene]
                    available=v is not None and n>0
                    summaries.append(dict(library=library,qc=qc_name,group=group,role=role,gene=gene,n=n,available=available,n_detected=int(np.count_nonzero(v[mask])) if available else '',fraction_detected=float(np.mean(v[mask]>0)) if available else '',mean_log1p_10k=float(np.mean(np.log1p(v[mask]*10000/total[mask]))) if available else '',aggregate_cpm=float(v[mask].sum()*1e6/total[mask].sum()) if available else ''))
            # Cell table carries no expression outcomes; supports independent selection audit.
            for j,b in enumerate(barcodes):
                per_cell.append(dict(library=library,barcode=b,qc=qc_name,umi=int(total[j]),n_genes=int(ngenes[j]),mito_fraction=float(mito[j]),qc_pass=bool(qc[j]),epi_features_pass=bool(epi[j]),colineage_flag=bool(mixed[j]),epithelial_candidate=bool(eligible_epi[j]),AT2_candidate=bool(groups['AT2_candidate'][j]),published_label=labels[j]))
        print(library,coverage[-2:],flush=True)
    differences=[]
    lookup={(r['library'],r['qc'],r['group'],r['gene']):r for r in summaries}
    for qc in cfg['qc']:
        for group in ['epithelial_candidate','AT2_candidate','AT2_published_intersection']:
            for role,gene in readouts:
                a=lookup[(cfg['primary_comparison'][0],qc,group,gene)]; b=lookup[(cfg['primary_comparison'][1],qc,group,gene)]
                limit=cfg['min_epithelial_for_comparison'] if group=='epithelial_candidate' else cfg['min_at2_for_comparison']
                valid=a['available'] and b['available'] and min(a['n'],b['n'])>=limit
                differences.append(dict(qc=qc,group=group,role=role,gene=gene,pam_n=a['n'],control_n=b['n'],coverage_pass=valid,mean_log1p_difference=a['mean_log1p_10k']-b['mean_log1p_10k'] if valid else '',detection_percentage_point_difference=100*(a['fraction_detected']-b['fraction_detected']) if valid else ''))
    for name,rows in [('coverage',coverage),('marker_summaries',summaries),('marker_differences',differences),('gene_coverage',genes),('published_annotation_crosswalk',cross),('annotation_coverage',annotation_coverage),('cell_selection',per_cell)]: write_table(tables/(name+'.tsv'),rows)
    inventory=[]; values=[]; stats=[]; stone=[]
    for sheet in book:
        # Read headings only here, never scan the large expression sheet as assay linkage evidence.
        header=[str(c.value) for row in sheet.iter_rows(min_row=1,max_row=min(5,sheet.max_row)) for c in row if c.value is not None]
        inventory.append(dict(sheet=sheet.title,rows=sheet.max_row,columns=sheet.max_column,header=' | '.join(header)[:600],scope='selected biochemical endpoint' if sheet.title in bio['sheets'] else 'inventory only'))
    for sheet_name in bio['sheets']:
        sheet=book[sheet_name]
        if sheet_name in ['Fig 6d','Fig 6e','Fig 6f']:
            for row in range(5,9):
                before,after=sheet.cell(row,2).value,sheet.cell(row,3).value
                if before is None or after is None: raise ValueError('Missing paired stone measurement')
                stone.append(dict(sheet=sheet_name,source_row=row,source_pair='row position; subject ID absent',diet=sheet.cell(2,1).value,D0_g=before,D56_g=after,D56_D0_ratio=after/before,percent_change=100*(after/before-1)))
            continue
        for col in range(3,9):
            genotype='Npt2b+/+' if col<6 else 'Npt2b-/-'
            diet=sheet.cell(5,col).value
            numbers=[]
            for row in range(6,sheet.max_row+1):
                v=sheet.cell(row,col).value
                if v is None: continue
                if not isinstance(v,(float,int)): raise ValueError('Unexpected assay value')
                numbers.append(v)
                values.append(dict(sheet=sheet_name,coordinate=sheet.cell(row,col).coordinate,endpoint=sheet.cell(2,1).value if sheet_name in ['Fig 6s','Fig 6t'] else sheet.cell(1,1).value,genotype=genotype,diet=diet,value=v,subject_id='not supplied',cross_assay_link='unresolved'))
            stats.append(dict(sheet=sheet_name,genotype=genotype,diet=diet,n=len(numbers),mean=float(np.mean(numbers)),median=float(np.median(numbers)),sd=float(np.std(numbers,ddof=1))))
    contrasts=[]
    for sn in [s for s in bio['sheets'] if s not in ['Fig 6d','Fig 6e','Fig 6f']]:
        for genotype in ['Npt2b+/+','Npt2b-/-']:
            a=next(s for s in stats if (s['sheet'],s['genotype'],s['diet'])==(sn,genotype,'LPD'))
            b=next(s for s in stats if (s['sheet'],s['genotype'],s['diet'])==(sn,genotype,'RD'))
            contrasts.append(dict(sheet=sn,genotype=genotype,LPD_n=a['n'],RD_n=b['n'],mean_LPD_minus_RD=a['mean']-b['mean'],mean_LPD_div_RD=a['mean']/b['mean'] if b['mean'] else '',inference='descriptive; separate assay, no cross-assay subject matching'))
    book.close()
    for name,rows in [('workbook_inventory',inventory),('biochemical_values',values),('biochemical_summary',stats),('biochemical_contrasts',contrasts),('stone_pairs',stone)]: write_table(tables/(name+'.tsv'),rows)
    inputs=[Path(__file__),RQ/'config/external_pilot_v2.json',RQ/'config/external_pilot_v1_annotation_addendum.json',RQ/'config/biochemical_audit_v1.json']+sorted(acq.glob('*.json'))+[ROOT/r['path'] for r in sources['source_files']+matrices]
    inputs=list(dict.fromkeys(inputs))
    record={'run':RUN,'completed_utc':datetime.datetime.now(datetime.timezone.utc).isoformat(),'inference':'descriptive external case context and separate source-assay reproduction; no new causal model','environment':{'python':platform.python_version(),'numpy':np.__version__,'scipy':scipy.__version__,'h5py':h5py.__version__,'openpyxl':openpyxl.__version__},'inputs':[dict(path=p.relative_to(ROOT).as_posix(),sha256=digest(p)) for p in inputs],'outputs':[dict(path=p.relative_to(ROOT).as_posix(),sha256=digest(p)) for p in sorted(tables.glob('*.tsv'))]}
    (paths['metadata']/'run_record.json').write_text(json.dumps(record,indent=2)+'\n',encoding='utf-8')
    print('Saved',len(record['outputs']),'tables; preserved units and descriptive scope.')

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--check',action='store_true');parser.add_argument('--tracked-only',action='store_true');a=parser.parse_args()
    if a.check: check(a.tracked_only)
    else: run()
